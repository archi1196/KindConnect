from flask import Flask, render_template, request, redirect, url_for, flash, send_file, session
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import sqlite3
import os
import hmac
from pathlib import Path
from datetime import date

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "kindconnect-local-development-key-change-before-deploy")
SITE_PASSWORD = os.environ.get("SITE_PASSWORD", "USER@12340")
DB_PATH = Path(__file__).with_name("community_welfare.db")

ACTIVITY_TYPES = [
    "Food Distribution (Annadaan)",
    "Tree Plantation",
    "Bird Feeding & Water",
    "Animal Welfare",
    "Blood Donation Awareness",
    "Community Survey",
    "Awareness Programme",
    "Other",
]

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


EXCEL_PATH = Path(__file__).with_name("KindConnect_Live_Records.xlsx")

def sync_excel():
    """Rewrite the local Excel workbook from the current database after every change."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    with get_db() as conn:
        activities = conn.execute("""
            SELECT id, title, activity_type, location, activity_date, volunteer, notes, status, created_at
            FROM activities ORDER BY activity_date DESC, id DESC
        """).fetchall()
        volunteers = conn.execute("""
            SELECT id, name, email, phone, interest, created_at
            FROM volunteers ORDER BY id DESC
        """).fetchall()

    workbook = Workbook()
    activity_sheet = workbook.active
    activity_sheet.title = "Activities"
    activity_sheet.append(["ID", "Activity Title", "Activity Type", "Location", "Activity Date",
                           "Volunteer / Team", "Notes", "Status", "Created At"])
    for row in activities:
        activity_sheet.append(list(row))

    volunteer_sheet = workbook.create_sheet("Volunteer Registrations")
    volunteer_sheet.append(["ID", "Full Name", "Email", "Phone", "Area of Interest", "Registered At"])
    for row in volunteers:
        volunteer_sheet.append(list(row))

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176B4D")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for column_cells in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_length + 2, 12), 42)

    workbook.save(EXCEL_PATH)

def init_db():
    with get_db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS activities (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                activity_type TEXT NOT NULL,
                location TEXT NOT NULL,
                activity_date TEXT NOT NULL,
                volunteer TEXT DEFAULT '',
                notes TEXT DEFAULT '',
                status TEXT NOT NULL DEFAULT 'Pending',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS volunteers (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                email TEXT NOT NULL,
                phone TEXT DEFAULT '',
                interest TEXT DEFAULT '',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(email)
            )
        """)
        count = conn.execute("SELECT COUNT(*) AS n FROM activities").fetchone()["n"]
        if count == 0:
            today = date.today()
            sample_activities = [
                ("Annadaan food packet distribution", "Food Distribution (Annadaan)", "Nanpura, Surat", today.isoformat(), "Volunteer Team", "DEMO RECORD — replace with actual activity details.", "Completed"),
                ("Tree sapling care", "Tree Plantation", "Adajan, Surat", today.isoformat(), "Green Team", "DEMO RECORD — check watering and plant health.", "Pending"),
                ("Bird water bowl check", "Bird Feeding & Water", "Athwa, Surat", today.isoformat(), "Community Volunteers", "DEMO RECORD — refill clean water as needed.", "Pending"),
                ("Animal welfare support", "Animal Welfare", "Vesu, Surat", today.isoformat(), "Care Team", "DEMO RECORD — sample only.", "In Progress"),
                ("Blood donation awareness drive", "Blood Donation Awareness", "Nanpura, Surat", today.isoformat(), "Awareness Team", "DEMO RECORD — sample only.", "Completed"),
                ("Community needs survey", "Community Survey", "Udhna, Surat", today.isoformat(), "Survey Team", "DEMO RECORD — sample only.", "Pending"),
                ("Public cleanliness awareness", "Awareness Programme", "Varachha, Surat", today.isoformat(), "Student Volunteers", "DEMO RECORD — sample only.", "In Progress"),
                ("Annadaan planning and packing", "Food Distribution (Annadaan)", "Surat", today.isoformat(), "Food Team", "DEMO RECORD — sample only.", "Completed"),
                ("Plantation follow-up visit", "Tree Plantation", "Pal, Surat", today.isoformat(), "Green Team", "DEMO RECORD — sample only.", "Pending"),
                ("Community water bowl refill", "Bird Feeding & Water", "Citylight, Surat", today.isoformat(), "Care Team", "DEMO RECORD — sample only.", "Completed"),
            ]
            conn.executemany("""
                INSERT INTO activities (title, activity_type, location, activity_date, volunteer, notes, status)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            """, sample_activities)

    sync_excel()

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        supplied = request.form.get("password", "")
        if hmac.compare_digest(supplied, SITE_PASSWORD):
            session["authenticated"] = True
            return redirect(url_for("index"))
        flash("Incorrect password. Please try again.", "error")
    return render_template("login.html")

@app.post("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

@app.before_request
def require_site_password():
    # Only the login page, static assets and health check are accessible before login.
    if request.endpoint in ("login", "static", "health"):
        return None
    if not session.get("authenticated"):
        return redirect(url_for("login", next=request.path))

@app.route("/")
def index():
    start_date = request.args.get("start_date", "").strip()
    end_date = request.args.get("end_date", "").strip()
    query = "SELECT * FROM activities WHERE 1=1"
    params = []
    if start_date:
        query += " AND activity_date >= ?"
        params.append(start_date)
    if end_date:
        query += " AND activity_date <= ?"
        params.append(end_date)
    query += " ORDER BY activity_date DESC, id DESC"
    with get_db() as conn:
        activities = conn.execute(query, params).fetchall()
        volunteers = conn.execute("SELECT * FROM volunteers ORDER BY id DESC LIMIT 8").fetchall()
        stats = {
            "activities": conn.execute("SELECT COUNT(*) FROM activities").fetchone()[0],
            "completed": conn.execute("SELECT COUNT(*) FROM activities WHERE status='Completed'").fetchone()[0],
            "pending": conn.execute("SELECT COUNT(*) FROM activities WHERE status='Pending'").fetchone()[0],
            "volunteers": conn.execute("SELECT COUNT(*) FROM volunteers").fetchone()[0],
        }
    return render_template("index.html", activities=activities, volunteers=volunteers,
                           stats=stats, activity_types=ACTIVITY_TYPES, today=date.today().isoformat(),
                           start_date=start_date, end_date=end_date)

@app.post("/activities/add")
def add_activity():
    title = request.form.get("title", "").strip()
    activity_type = request.form.get("activity_type", "").strip()
    location = request.form.get("location", "").strip()
    activity_date = request.form.get("activity_date", "").strip()
    volunteer = request.form.get("volunteer", "").strip()
    notes = request.form.get("notes", "").strip()
    if not all([title, activity_type, location, activity_date]):
        flash("Please complete all required activity fields.", "error")
        return redirect(url_for("index") + "#activities")
    if activity_type not in ACTIVITY_TYPES:
        flash("Please choose a valid activity type.", "error")
        return redirect(url_for("index") + "#activities")
    with get_db() as conn:
        conn.execute("""
            INSERT INTO activities (title, activity_type, location, activity_date, volunteer, notes)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (title, activity_type, location, activity_date, volunteer, notes))
    sync_excel()
    flash("Activity added successfully.", "success")
    return redirect(url_for("index") + "#activities")

@app.post("/activities/<int:activity_id>/status")
def update_status(activity_id):
    status = request.form.get("status", "")
    if status not in ("Pending", "In Progress", "Completed"):
        flash("Invalid activity status.", "error")
    else:
        with get_db() as conn:
            conn.execute("UPDATE activities SET status=? WHERE id=?", (status, activity_id))
        sync_excel()
        flash("Activity status updated.", "success")
    return redirect(url_for("index") + "#activities")

@app.post("/activities/<int:activity_id>/delete")
def delete_activity(activity_id):
    with get_db() as conn:
        conn.execute("DELETE FROM activities WHERE id=?", (activity_id,))
    sync_excel()
    flash("Activity deleted.", "success")
    return redirect(url_for("index") + "#activities")

@app.post("/volunteers/add")
def add_volunteer():
    name = request.form.get("name", "").strip()
    email = request.form.get("email", "").strip().lower()
    phone = request.form.get("phone", "").strip()
    interest = request.form.get("interest", "").strip()
    if not name or not email or "@" not in email:
        flash("Please enter a name and a valid email address.", "error")
        return redirect(url_for("index") + "#volunteers")
    try:
        with get_db() as conn:
            conn.execute("INSERT INTO volunteers (name, email, phone, interest) VALUES (?, ?, ?, ?)",
                         (name, email, phone, interest))
        sync_excel()
        flash("Volunteer registered successfully. Your details have been saved to the Excel records.", "success")
    except sqlite3.IntegrityError:
        flash("That email is already registered.", "error")
    return redirect(url_for("index") + "#volunteers")


@app.get("/export/excel")
def export_excel():
    """Download an Excel workbook; activity rows respect the selected date range."""
    start_date = request.args.get("start_date", "").strip()
    end_date = request.args.get("end_date", "").strip()
    if start_date and end_date and start_date > end_date:
        flash("Start date must be on or before end date.", "error")
        return redirect(url_for("index") + "#activities")

    with get_db() as conn:
        activity_sql = "SELECT id, title, activity_type, location, activity_date, volunteer, notes, status, created_at FROM activities WHERE 1=1"
        activity_params = []
        if start_date:
            activity_sql += " AND activity_date >= ?"
            activity_params.append(start_date)
        if end_date:
            activity_sql += " AND activity_date <= ?"
            activity_params.append(end_date)
        activity_sql += " ORDER BY activity_date DESC, id DESC"
        activities = conn.execute(activity_sql, activity_params).fetchall()
        volunteers = conn.execute("SELECT id, name, email, phone, interest, created_at FROM volunteers ORDER BY id DESC").fetchall()

    workbook = Workbook()
    activity_sheet = workbook.active
    activity_sheet.title = "Activities"
    activity_sheet.append(["ID", "Activity Title", "Activity Type", "Location", "Activity Date", "Volunteer / Team", "Notes", "Status", "Created At"])
    for row in activities:
        activity_sheet.append(list(row))
    volunteer_sheet = workbook.create_sheet("Volunteer Registrations")
    volunteer_sheet.append(["ID", "Full Name", "Email", "Phone", "Area of Interest", "Registered At"])
    for row in volunteers:
        volunteer_sheet.append(list(row))

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173F32")
            cell.alignment = Alignment(vertical="center")
        for column_cells in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_length + 2, 12), 42)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    suffix = "_".join(x for x in (start_date, end_date) if x) or "all_dates"
    return send_file(output, as_attachment=True,
                     download_name=f"KindConnect_Records_{suffix}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/health")
def health():
    return {"status": "ok", "application": "Digital Community Welfare Management System"}

init_db()

if __name__ == "__main__":
    print("Digital Community Welfare Management System is running at http://127.0.0.1:5000")
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
